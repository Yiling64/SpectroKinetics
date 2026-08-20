import datetime
import io
import os
import re
import zipfile
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
import pandas as pd
import streamlit as st

# ================= 頁面全域設定 =================
st.set_page_config(
    page_title="SpectroKinetics Pro",
    layout="wide",
    page_icon="🔬",
    initial_sidebar_state="expanded",
)
plt.rcParams["font.sans-serif"] = "Arial"
plt.rcParams["font.family"] = "sans-serif"
plt.rcParams["axes.unicode_minus"] = False

if "parsed_data" not in st.session_state:
    st.session_state.parsed_data = {}
if "df_meta" not in st.session_state:
    st.session_state.df_meta = pd.DataFrame()
if "global_unit" not in st.session_state:
    st.session_state.global_unit = "mg/ml"
if "editor_version" not in st.session_state:
    st.session_state.editor_version = 0
if "selected_comp" not in st.session_state:
    st.session_state.selected_comp = None
if "history_stack" not in st.session_state:
    st.session_state.history_stack = []

# ================= 專業軟體 UI 樣式系統 =================
CUSTOM_UI_CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');
html, body, [class*="css"] {
    font-family: 'Inter', -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
    color: #1e293b;
    background-color: #f8fafc;
}

#MainMenu, footer {visibility: hidden;}
header {
    background: transparent !important;
}

.block-container {
    padding-top: 2.2rem !important;
    padding-bottom: 1.5rem !important;
    max-width: 98% !important;
}

/* 側邊控制台 */
section[data-testid="stSidebar"] {
    background-color: #f8fafc !important;
    color: #1e293b !important;
    border-right: 1px solid #e2e8f0;
}
section[data-testid="stSidebar"] h1, 
section[data-testid="stSidebar"] h2, 
section[data-testid="stSidebar"] h3,
section[data-testid="stSidebar"] label,
section[data-testid="stSidebar"] p,
section[data-testid="stSidebar"] span {
    color: #0f172a !important;
    font-weight: 500;
}
section[data-testid="stSidebar"] input {
    background-color: #ffffff !important;
    color: #0f172a !important;
    border: 1px solid #cbd5e1 !important;
    border-radius: 6px !important;
}

/* 檔案上傳區塊 */
section[data-testid="stSidebar"] [data-testid="stFileUploader"] {
    background-color: #ffffff !important;
    border-radius: 8px !important;
    padding: 6px !important;
    border: 1px dashed #94a3b8 !important;
}
section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] {
    background-color: #f8fafc !important;
}
section[data-testid="stSidebar"] [data-testid="stFileUploader"] small,
section[data-testid="stSidebar"] [data-testid="stFileUploader"] span,
section[data-testid="stSidebar"] [data-testid="stFileUploader"] button {
    color: #334155 !important;
}

/* 頂部橫幅 */
.app-header {
    display: flex;
    justify-content: space-between;
    align-items: center;
    background: #ffffff;
    padding: 12px 20px;
    border-radius: 8px;
    border: 1px solid #e2e8f0;
    box-shadow: 0 1px 3px rgba(0,0,0,0.04);
    margin-bottom: 1rem;
}
.app-title {
    font-size: 1.25rem;
    font-weight: 700;
    color: #0f172a;
    display: flex;
    align-items: center;
    gap: 8px;
}
.app-badge {
    font-size: 0.75rem;
    font-weight: 700;
    background: #e0f2fe;
    color: #0284c7;
    padding: 3px 10px;
    border-radius: 4px;
    border: 1px solid #bae6fd;
    text-transform: uppercase;
    letter-spacing: 0.5px;
}

/* 卡片容器 */
.workspace-card {
    background: #ffffff;
    border: 1px solid #e2e8f0;
    border-radius: 8px;
    padding: 16px;
    margin-bottom: 12px;
    box-shadow: 0 1px 2px rgba(0,0,0,0.02);
}
.card-header {
    font-size: 0.95rem;
    font-weight: 700;
    color: #1e293b;
    border-bottom: 1px solid #f1f5f9;
    padding-bottom: 8px;
    margin-bottom: 12px;
    display: flex;
    justify-content: space-between;
    align-items: center;
}

/* 色彩標籤與區塊 */
.tag-badge {
    display: inline-block;
    padding: 2px 8px;
    border-radius: 4px;
    font-size: 11px;
    font-weight: 700;
    letter-spacing: 0.3px;
}
.tag-date {
    background: #f1f5f9;
    color: #475569;
    border: 1px solid #cbd5e1;
}
.tag-comp {
    background: #f0fdf4;
    color: #166534;
    border: 1px solid #bbf7d0;
}
.tag-mode {
    background: #faf5ff;
    color: #6b21a8;
    border: 1px solid #e9d5ff;
}

.date-divider {
    background: #f8fafc;
    border-left: 4px solid #0284c7;
    border-top: 1px solid #e2e8f0;
    border-right: 1px solid #e2e8f0;
    border-bottom: 1px solid #e2e8f0;
    border-radius: 0 6px 6px 0;
    padding: 8px 14px;
    margin: 12px 0 8px 0;
    font-weight: 700;
    color: #0f172a;
    font-size: 13px;
    display: flex;
    justify-content: space-between;
    align-items: center;
}

/* Tab 樣式 */
.stTabs [data-baseweb="tab-list"] {
    gap: 6px;
    background-color: #e2e8f0;
    padding: 4px;
    border-radius: 8px;
}
.stTabs [data-baseweb="tab"] {
    height: 36px;
    border-radius: 6px;
    font-size: 13px;
    font-weight: 600;
    color: #475569;
    background: transparent;
    padding: 0 14px;
}
.stTabs [aria-selected="true"] {
    background-color: #ffffff !important;
    color: #0f172a !important;
    box-shadow: 0 1px 3px rgba(0,0,0,0.08);
}
</style>
"""
st.markdown(CUSTOM_UI_CSS, unsafe_allow_html=True)


# ================= 輔助函式 =================
def build_sample_label(comp, dose, unit, suffix):
    comp_str = str(comp).strip()
    if any(
        k in comp_str.lower()
        for k in ["control", "basal", "water+water", "water+reagent"]
    ):
        return comp_str
    try:
        d_num = float(dose)
        dose_str = f"{d_num:g}" if d_num != 0 else ""
    except (ValueError, TypeError):
        dose_str = str(dose).strip() if str(dose).strip() != "0" else ""

    unit_str = str(unit).strip() if str(unit).strip() != "-" else ""
    parts = [comp_str]
    if dose_str and unit_str:
        parts.append(f"{dose_str} {unit_str}")
    elif dose_str:
        parts.append(dose_str)
    elif unit_str:
        parts.append(unit_str)

    main_label = " ".join(parts).strip()
    suf = str(suffix).strip() if pd.notna(suffix) else ""
    if suf:
        suf_clean = re.sub(r"[\(\)]", "", suf).strip()
        if suf_clean:
            main_label = f"{main_label} ({suf_clean})"
    return main_label


def smart_parse_sample(sample_name, default_unit="mg/ml"):
    s = sample_name.strip()
    s_lower = s.lower()
    if any(k in s_lower for k in ["water+water", "w+w", "water + water"]):
        return "Blank (Water+Water)", "0", "-", ""
    if any(
        k in s_lower
        for k in [
            "water+reagent",
            "w+reagent",
            "water + reagent",
            "water+試劑",
            "水+試劑",
            "水+水",
        ]
    ):
        return "Blank (Water+Reagent)", "0", "-", ""
    if "control" in s_lower or "ctrl" in s_lower:
        return "Control", "0", "-", ""
    if "basal" in s_lower or "basel" in s_lower:
        return "Basal", "0", "-", ""

    suffix_found = ""
    suf_match = re.search(
        r"(?i)\s*(alone|單獨|单独|only|blank|water|\+水|\+ water)\s*", s
    )
    if suf_match:
        suffix_found = suf_match.group(1).strip("+ ")

    clean_s = re.sub(
        r"(?i)\s*(alone|單獨|单独|only|blank|water|\+水|\+ water)\s*", "", s
    ).strip()
    dose_unit_pattern = (
        r"(.*?)(?:[_\s]+)?([0-9]+\.?[0-9]*)\s*(mM|μM|uM|μg/ml|ug/ml|mg/ml|nM|M|%)?$"
    )
    match = re.search(dose_unit_pattern, clean_s, re.IGNORECASE)

    if match:
        raw_comp = match.group(1).strip(" _-")
        dose = match.group(2)
        unit = match.group(3) if match.group(3) else ""
    else:
        raw_comp = clean_s
        dose = "0"
        unit = ""

    compound = raw_comp if raw_comp else "Unknown"
    if not unit:
        unit = "mM" if compound.lower().startswith("pro") else default_unit
    return compound, dose, unit, suffix_found


def parse_txt_files(uploaded_files, def_unit):
    parsed_dict = {}
    meta_list = []
    global_item_counter = 0

    for file in uploaded_files:
        content = file.read().decode("utf-8", errors="ignore")
        filename = file.name
        date_match = re.search(
            r"Run Date:\t'[\d:]+, (\d{2}/\d{2}/\d{4})", content
        )
        file_date = (
            date_match.group(1).replace("/", "-") if date_match else "Unknown"
        )

        blocks = content.split("Sample:\t'")
        ch_counter = 0

        for block in blocks:
            if not block.strip():
                continue
            lines = block.strip().split("\n")
            run_time = "99:99:99"
            for line in lines[:10]:
                match = re.search(r"Run Date:\t'(\d{2}:\d{2}:\d{2})", line)
                if match:
                    run_time = match.group(1)
                    break

            sample_name = lines[0].split("\t")[0].strip().strip("'").strip()
            if sample_name == "1" or "Data Points" not in block:
                continue

            ch_counter += 1
            global_item_counter += 1
            data_part = block.split("Data Points")[1]
            times, values = [], []
            for d_line in data_part.strip().split("\n"):
                parts = d_line.split()
                if len(parts) == 2:
                    try:
                        times.append(float(parts[0]))
                        values.append(float(parts[1]))
                    except ValueError:
                        pass

            if times:
                unique_id = f"{file_date}_{run_time}_{sample_name}_ch{ch_counter}_{global_item_counter}_{filename}"
                df_points = pd.DataFrame({"Time": times, "Abs": values})
                parsed_dict[unique_id] = {"df": df_points, "ch_idx": ch_counter}
                comp, dose, unit, suffix_val = smart_parse_sample(
                    sample_name, def_unit
                )

                meta_list.append(
                    {
                        "選取": True,
                        "UID": unique_id,
                        "日期 (Date)": file_date,
                        "測量時間": run_time,
                        "化合物 (Compound)": comp,
                        "樣品名稱 (Sample)": sample_name,
                        "劑量 (Dose)": dose,
                        "單位 (Unit)": unit,
                        "後綴 (Suffix)": suffix_val,
                        "來源檔名": filename,
                    }
                )

    canonical_names = {}
    for item in meta_list:
        comp = item["化合物 (Compound)"]
        if any(
            k in comp.lower()
            for k in ["control", "basal", "water+water", "water+reagent"]
        ):
            continue
        low_k = comp.lower()
        if low_k not in canonical_names:
            canonical_names[low_k] = comp
        else:
            if comp[0].isupper() and not canonical_names[low_k][0].isupper():
                canonical_names[low_k] = comp

    for item in meta_list:
        low_k = item["化合物 (Compound)"].lower()
        if low_k in canonical_names:
            item["化合物 (Compound)"] = canonical_names[low_k]

    df_res = pd.DataFrame(meta_list)
    if not df_res.empty:
        df_res["_dt_sort"] = pd.to_datetime(
            df_res["日期 (Date)"], format="%d-%m-%Y", errors="coerce"
        )
        df_res = (
            df_res.sort_values(by=["_dt_sort", "測量時間"])
            .drop(columns=["_dt_sort"])
            .reset_index(drop=True)
        )

    return parsed_dict, df_res


# ================= 側邊欄控制面板 =================
st.sidebar.markdown(
    """
    <div style='padding: 6px 0 12px 0;'>
        <div style='font-size: 1.05rem; font-weight: 700; color: #0284c7; letter-spacing: 0.5px;'>[ EXPERIMENT CONFIG ]</div>
        <div style='font-size: 0.8rem; color: #64748b;'>實驗條件與反應時序設定</div>
    </div>
    """,
    unsafe_allow_html=True,
)

assay_type = st.sidebar.radio(
    "實驗類型 (Assay)", ["Superoxide Generation", "Elastase Release"]
)
is_superoxide = "Superoxide" in assay_type

st.sidebar.markdown("---")
st.sidebar.markdown(
    "<div style='font-size: 11px; font-weight: 700; color: #475569;"
    " letter-spacing: 0.5px;'>METADATA CONFIGURATION</div>",
    unsafe_allow_html=True,
)
meta_cell = st.sidebar.text_input("Cell", value="6 × 10⁵ cells/ml")
meta_stim = st.sidebar.text_input("Stimulant", value="fMLP 10⁻⁷ M")
meta_amp = st.sidebar.text_input(
    "Amplifier",
    value="Cyto B 1 μg/ml" if is_superoxide else "Cyto B 0.5 μg/ml",
)
meta_sub = st.sidebar.text_input(
    "Substrate",
    value=(
        "Cytochrome C 0.5 mg/ml" if is_superoxide else "Substrate 0.1 μM"
    ),
)

st.sidebar.markdown("---")
st.sidebar.markdown(
    "<div style='font-size: 11px; font-weight: 700; color: #475569;"
    " letter-spacing: 0.5px;'>TIMING PARAMETERS (SEC)</div>",
    unsafe_allow_html=True,
)
t_start = st.sidebar.number_input("Start", value=420)
t_interval = st.sidebar.number_input("Interval", value=15)
t_react = st.sidebar.number_input("Reaction", value=600)

st.sidebar.markdown("---")
st.sidebar.markdown(
    "<div style='font-size: 11px; font-weight: 700; color: #475569;"
    " letter-spacing: 0.5px;'>Y-AXIS SCALE CONFIG (O.D.)</div>",
    unsafe_allow_html=True,
)
y_min = st.sidebar.number_input(
    "Y 軸最小值 (Y-min)", value=0.0, step=0.1, format="%.2f"
)
y_max = st.sidebar.number_input(
    "Y 軸最大值 (Y-max)",
    value=1.2 if is_superoxide else 0.8,
    step=0.1,
    format="%.2f",
)

st.sidebar.markdown("---")
st.sidebar.markdown(
    "<div style='font-size: 11px; font-weight: 700; color: #475569;"
    " letter-spacing: 0.5px;'>RAW DATA IMPORT</div>",
    unsafe_allow_html=True,
)
uploaded_files = st.sidebar.file_uploader(
    "選擇檔案",
    accept_multiple_files=True,
    type=["txt"],
    label_visibility="collapsed",
)
if st.sidebar.button("解析數據檔", use_container_width=True):
    if uploaded_files:
        st.session_state.parsed_data, st.session_state.df_meta = parse_txt_files(
            uploaded_files, st.session_state.global_unit
        )
        st.session_state.editor_version += 1
        st.session_state.history_stack = []
        st.sidebar.success("檔案解析成功")
    else:
        st.sidebar.warning("請先選取 TXT 檔案")

# ================= 頂部橫幅 =================
header_left, header_right = st.columns([3, 1])
with header_left:
    st.markdown(
        f"""
        <div class="app-title">
            <span>SpectroKinetics Workspace</span>
            <span class="app-badge">{assay_type}</span>
        </div>
        """,
        unsafe_allow_html=True,
    )
with header_right:
    if st.session_state.history_stack:
        if st.button("復原上一步刪除 [Undo]", use_container_width=True):
            last_meta, last_parsed = st.session_state.history_stack.pop()
            st.session_state.df_meta = last_meta.copy()
            st.session_state.parsed_data = last_parsed.copy()
            st.session_state.editor_version += 1
            st.rerun()

# ================= 主工作區 =================
if not st.session_state.df_meta.empty:
    unit_options = ["μM", "mM", "μg/ml", "mg/ml", "nM", "M", "%", "-"]
    selected_df = st.session_state.df_meta[
        st.session_state.df_meta["選取"] == True
    ]

    compounds_all = sorted(
        st.session_state.df_meta["化合物 (Compound)"].unique(),
        key=lambda x: (
            0
            if any(
                k in x.lower()
                for k in ["basal", "control", "water+water", "water+reagent"]
            )
            else 1,
            x.lower(),
        ),
    )

    if (
        st.session_state.selected_comp not in compounds_all
        and len(compounds_all) > 0
    ):
        st.session_state.selected_comp = compounds_all[0]

    tab_data, tab_excel, tab_plot = st.tabs(
        [
            "  樣品檢視與對照組  ",
            "  Excel 報表結算  ",
            "  動力學曲線繪製  ",
        ]
    )

    # -------------------------------------------------------------
    # 分頁 1: 左側選單 + 右側依日期分區滾動
    # -------------------------------------------------------------
    with tab_data:
        col_nav, col_detail = st.columns([1.1, 3.1])

        with col_nav:
            st.markdown(
                """
                <div class="workspace-card" style="padding: 12px;">
                    <div class="card-header" style="margin-bottom: 8px;">樣品清單</div>
                """,
                unsafe_allow_html=True,
            )

            for comp in compounds_all:
                cnt = len(
                    st.session_state.df_meta[
                        st.session_state.df_meta["化合物 (Compound)"] == comp
                    ]
                )
                is_active = comp == st.session_state.selected_comp

                r_c1, r_c2 = st.columns([4, 1])
                with r_c1:
                    btn_type = "primary" if is_active else "secondary"
                    if st.button(
                        f"{comp} ({cnt})",
                        key=f"nav_btn_{comp}",
                        use_container_width=True,
                        type=btn_type,
                    ):
                        st.session_state.selected_comp = comp
                        st.rerun()

                with r_c2:
                    if st.button("✕", key=f"del_x_{comp}", help=f"刪除 {comp}"):
                        st.session_state.history_stack.append(
                            (
                                st.session_state.df_meta.copy(),
                                st.session_state.parsed_data.copy(),
                            )
                        )

                        uids_to_remove = st.session_state.df_meta[
                            st.session_state.df_meta["化合物 (Compound)"]
                            == comp
                        ]["UID"].tolist()
                        for uid in uids_to_remove:
                            st.session_state.parsed_data.pop(uid, None)

                        st.session_state.df_meta = st.session_state.df_meta[
                            st.session_state.df_meta["化合物 (Compound)"]
                            != comp
                        ].reset_index(drop=True)

                        rem_comps = sorted(
                            st.session_state.df_meta[
                                "化合物 (Compound)"
                            ].unique()
                        )
                        st.session_state.selected_comp = (
                            rem_comps[0] if rem_comps else None
                        )
                        st.session_state.editor_version += 1
                        st.rerun()

            st.markdown("---")
            st.caption("全域單位快速套用")
            sel_g_unit = st.selectbox(
                "套用單位",
                unit_options,
                index=unit_options.index(st.session_state.global_unit)
                if st.session_state.global_unit in unit_options
                else 3,
                label_visibility="collapsed",
            )
            if st.button("一鍵套用全部", use_container_width=True):
                st.session_state.global_unit = sel_g_unit
                mask = ~st.session_state.df_meta["化合物 (Compound)"].str.lower().isin(
                    [
                        "control",
                        "basal",
                        "blank (water+water)",
                        "blank (water+reagent)",
                    ]
                )
                st.session_state.df_meta.loc[mask, "單位 (Unit)"] = sel_g_unit
                st.session_state.editor_version += 1
                st.rerun()
            st.markdown("</div>", unsafe_allow_html=True)

        with col_detail:
            current_comp = st.session_state.selected_comp
            if current_comp:
                is_ctrl_comp = any(
                    k in current_comp.lower()
                    for k in [
                        "basal",
                        "control",
                        "water+water",
                        "water+reagent",
                    ]
                )

                st.markdown(
                    f"""
                    <div class="workspace-card">
                        <div class="card-header">
                            <span>當前檢視化合物：<span class="tag-badge tag-comp">{current_comp}</span></span>
                        </div>
                    """,
                    unsafe_allow_html=True,
                )

                h_col1, h_col2 = st.columns([2.2, 2.0])
                with h_col1:
                    if not is_ctrl_comp:
                        cur_comp_unit = st.session_state.df_meta[
                            st.session_state.df_meta["化合物 (Compound)"]
                            == current_comp
                        ]["單位 (Unit)"].values[0]
                        def_idx = (
                            unit_options.index(cur_comp_unit)
                            if cur_comp_unit in unit_options
                            else 3
                        )

                        def make_change_handler(target_comp):
                            def handler():
                                new_u = st.session_state[f"u_sel_{target_comp}"]
                                mask = (
                                    st.session_state.df_meta[
                                        "化合物 (Compound)"
                                    ]
                                    == target_comp
                                )
                                st.session_state.df_meta.loc[
                                    mask, "單位 (Unit)"
                                ] = new_u
                                st.session_state.editor_version += 1

                            return handler

                        st.selectbox(
                            f"變更 {current_comp} 單位:",
                            unit_options,
                            index=def_idx,
                            key=f"u_sel_{current_comp}",
                            on_change=make_change_handler(current_comp),
                        )

                with h_col2:
                    st.write("")
                    st.write("")
                    c_sub1, c_sub2 = st.columns(2)
                    with c_sub1:
                        if st.button(
                            "全選此藥物",
                            key=f"sel_all_{current_comp}",
                            use_container_width=True,
                        ):
                            mask = (
                                st.session_state.df_meta["化合物 (Compound)"]
                                == current_comp
                            )
                            st.session_state.df_meta.loc[mask, "選取"] = True
                            st.session_state.editor_version += 1
                            st.rerun()
                    with c_sub2:
                        if st.button(
                            "清除此藥物",
                            key=f"clr_all_{current_comp}",
                            use_container_width=True,
                        ):
                            mask = (
                                st.session_state.df_meta["化合物 (Compound)"]
                                == current_comp
                            )
                            st.session_state.df_meta.loc[mask, "選取"] = False
                            st.session_state.editor_version += 1
                            st.rerun()

                st.markdown("</div>", unsafe_allow_html=True)

                comp_df_all = st.session_state.df_meta[
                    st.session_state.df_meta["化合物 (Compound)"]
                    == current_comp
                ]
                comp_unique_dates = sorted(
                    comp_df_all["日期 (Date)"].unique(),
                    key=lambda x: pd.to_datetime(
                        x, format="%d-%m-%Y", errors="coerce"
                    ),
                )

                for d in comp_unique_dates:
                    date_mask = (
                        st.session_state.df_meta["化合物 (Compound)"]
                        == current_comp
                    ) & (st.session_state.df_meta["日期 (Date)"] == d)
                    date_comp_df = st.session_state.df_meta[date_mask]

                    st.markdown(
                        f"""
                        <div class="workspace-card">
                            <div class="card-header" style="color: #0369a1;">
                                <span>實驗日期：<span class="tag-badge tag-date">{d}</span> ({len(date_comp_df)} 筆樣品)</span>
                            </div>
                        """,
                        unsafe_allow_html=True,
                    )

                    d_c1, d_c2, _ = st.columns([1.2, 1.2, 5])
                    with d_c1:
                        if st.button(
                            f"全選此日",
                            key=f"sel_{current_comp}_{d}",
                            use_container_width=True,
                        ):
                            st.session_state.df_meta.loc[date_mask, "選取"] = (
                                True
                            )
                            st.session_state.editor_version += 1
                            st.rerun()
                    with d_c2:
                        if st.button(
                            f"清除此日",
                            key=f"clr_{current_comp}_{d}",
                            use_container_width=True,
                        ):
                            st.session_state.df_meta.loc[date_mask, "選取"] = (
                                False
                            )
                            st.session_state.editor_version += 1
                            st.rerun()

                    edited_comp_date = st.data_editor(
                        date_comp_df,
                        column_config={
                            "選取": st.column_config.CheckboxColumn(
                                required=True
                            ),
                            "UID": None,
                            "日期 (Date)": st.column_config.TextColumn(
                                "日期 (Date)"
                            ),
                            "測量時間": st.column_config.TextColumn(
                                "測量時間"
                            ),
                            "化合物 (Compound)": st.column_config.TextColumn(
                                "化合物名稱", required=True
                            ),
                            "樣品名稱 (Sample)": st.column_config.TextColumn(
                                "樣品名稱"
                            ),
                            "劑量 (Dose)": st.column_config.TextColumn(
                                "劑量"
                            ),
                            "單位 (Unit)": st.column_config.SelectboxColumn(
                                "單位", options=unit_options, required=True
                            ),
                            "後綴 (Suffix)": st.column_config.TextColumn(
                                "標籤備註 (後綴)"
                            ),
                            "來源檔名": st.column_config.TextColumn(
                                "來源檔名"
                            ),
                        },
                        disabled=["日期 (Date)", "測量時間", "來源檔名"],
                        use_container_width=True,
                        hide_index=True,
                        key=f"editor_{current_comp}_{d}_v{st.session_state.editor_version}",
                    )
                    st.session_state.df_meta.update(edited_comp_date)
                    st.markdown("</div>", unsafe_allow_html=True)

        # 每日對照組基準綁定
        st.markdown(
            """
            <div class="workspace-card">
                <div class="card-header">每日對照組基準綁定 (Anchoring)</div>
            """,
            unsafe_allow_html=True,
        )
        anchor_settings = {}
        selected_df_curr = st.session_state.df_meta[
            st.session_state.df_meta["選取"] == True
        ]
        unique_active_dates = sorted(
            selected_df_curr["日期 (Date)"].unique(),
            key=lambda x: pd.to_datetime(
                x, format="%d-%m-%Y", errors="coerce"
            ),
        )
        cols = st.columns(
            len(unique_active_dates) if len(unique_active_dates) > 0 else 1
        )

        for idx, date in enumerate(unique_active_dates):
            with cols[idx]:
                st.markdown(f"**日期: `{date}`**")
                date_samples = selected_df_curr[
                    selected_df_curr["日期 (Date)"] == date
                ]
                options = date_samples["UID"].tolist()
                display_names = [
                    f"{r['樣品名稱 (Sample)']} ({r['測量時間']})"
                    for _, r in date_samples.iterrows()
                ]
                format_dict = dict(zip(options, display_names))

                basal_def = next(
                    (
                        u
                        for u, n in zip(options, display_names)
                        if "basal" in n.lower()
                    ),
                    options[0] if options else None,
                )
                ctrl_def = next(
                    (
                        u
                        for u, n in zip(options, display_names)
                        if "control" in n.lower()
                    ),
                    options[0] if options else None,
                )

                anchor_settings[date] = {}
                if not is_superoxide:
                    anchor_settings[date]["Basal"] = st.selectbox(
                        f"Basal 基準",
                        options,
                        index=options.index(basal_def)
                        if basal_def in options
                        else 0,
                        format_func=lambda x: format_dict.get(x, x),
                        key=f"b_anch_{date}",
                    )
                anchor_settings[date]["Control"] = st.selectbox(
                    f"Control 基準",
                    options,
                    index=options.index(ctrl_def)
                    if ctrl_def in options
                    else 0,
                    format_func=lambda x: format_dict.get(x, x),
                    key=f"c_anch_{date}",
                )
        st.markdown("</div>", unsafe_allow_html=True)

    # -------------------------------------------------------------
    # 分頁 2: 結算與匯出 Excel
    # -------------------------------------------------------------
    selected_df = st.session_state.df_meta[
        st.session_state.df_meta["選取"] == True
    ]
    with tab_excel:
        st.markdown(
            """
            <div class="workspace-card">
                <div class="card-header">Excel 試算表結算與匯出</div>
            """,
            unsafe_allow_html=True,
        )
        exp_c1, exp_c2 = st.columns([3, 2])
        with exp_c1:
            out_filename = st.text_input(
                "輸出檔案名稱",
                value=f"{'Superoxide' if is_superoxide else 'Elastase'}_Data_Summary.xlsx",
            )
        with exp_c2:
            existing_excel = st.file_uploader(
                "追加至既有 Excel 試算表 (選填)", type=["xlsx"]
            )

        if st.button("開始計算並建置 Excel 總表", use_container_width=True):
            if selected_df.empty:
                st.warning("請先在第一頁勾選欲結算的樣品！")
            else:
                calc_dict = {}
                for _, row in selected_df.iterrows():
                    uid = row["UID"]
                    df_points = st.session_state.parsed_data[uid]["df"]
                    ch_idx = st.session_state.parsed_data[uid]["ch_idx"]

                    target_min = t_start + (ch_idx - 1) * t_interval
                    target_max = target_min + t_react

                    val_min = df_points.iloc[
                        (df_points["Time"] - target_min).abs().idxmin()
                    ]["Abs"]
                    val_max = df_points.iloc[
                        (df_points["Time"] - target_max).abs().idxmin()
                    ]["Abs"]
                    calc_dict[uid] = {
                        "min": val_min,
                        "max": val_max,
                        "diff": val_max - val_min,
                    }

                compounds = [
                    c
                    for c in selected_df["化合物 (Compound)"].unique()
                    if not any(
                        k in c.lower()
                        for k in [
                            "control",
                            "basal",
                            "water+water",
                            "water+reagent",
                        ]
                    )
                ]

                if existing_excel:
                    wb = openpyxl.load_workbook(existing_excel)
                else:
                    wb = openpyxl.Workbook()
                    if "Sheet" in wb.sheetnames and len(wb.sheetnames) == 1:
                        wb.remove(wb["Sheet"])

                tnr_font = Font(name="Times New Roman", size=12)
                tnr_bold = Font(name="Times New Roman", size=12, bold=True)
                gray_fill = PatternFill(
                    start_color="E8E8E8", end_color="E8E8E8", fill_type="solid"
                )

                for comp in compounds:
                    clean_sheet_title = re.sub(
                        r"[\:\\/\?\*\[\]]", "_", str(comp)
                    )[:31].strip()
                    if not clean_sheet_title:
                        clean_sheet_title = "Sheet"

                    comp_df = selected_df[
                        selected_df["化合物 (Compound)"] == comp
                    ]
                    comp_dates = sorted(
                        comp_df["日期 (Date)"].unique(),
                        key=lambda x: pd.to_datetime(
                            x, format="%d-%m-%Y", errors="coerce"
                        ),
                    )
                    processed_rows = []

                    for d in comp_dates:
                        if d not in anchor_settings:
                            continue
                        c_uid = anchor_settings[d]["Control"]
                        c_info = selected_df[
                            selected_df["UID"] == c_uid
                        ].iloc[0]
                        c_diff = calc_dict[c_uid]["diff"]

                        if not is_superoxide:
                            b_uid = anchor_settings[d]["Basal"]
                            b_info = selected_df[
                                selected_df["UID"] == b_uid
                            ].iloc[0]
                            b_diff = calc_dict[b_uid]["diff"]
                            processed_rows.append(
                                [
                                    d,
                                    b_info["來源檔名"],
                                    "basal",
                                    round(calc_dict[b_uid]["min"], 4),
                                    round(calc_dict[b_uid]["max"], 4),
                                    round(b_diff, 4),
                                    0.0,
                                    "0.00%",
                                    "-",
                                ]
                            )

                        if is_superoxide:
                            processed_rows.append(
                                [
                                    d,
                                    c_info["來源檔名"],
                                    "control",
                                    round(calc_dict[c_uid]["min"], 4),
                                    round(calc_dict[c_uid]["max"], 4),
                                    round(c_diff, 4),
                                    round(c_diff * 47.4, 4),
                                    "100.00%",
                                    "0.00%",
                                ]
                            )
                        else:
                            c_net = c_diff - b_diff
                            processed_rows.append(
                                [
                                    d,
                                    c_info["來源檔名"],
                                    "control",
                                    round(calc_dict[c_uid]["min"], 4),
                                    round(calc_dict[c_uid]["max"], 4),
                                    round(c_diff, 4),
                                    round(c_net, 4),
                                    "100.00%",
                                    "0.00%",
                                ]
                            )

                        date_comp_df = comp_df[comp_df["日期 (Date)"] == d]
                        date_drug_rows = []

                        for _, row in date_comp_df.iterrows():
                            uid = row["UID"]
                            s_diff = calc_dict[uid]["diff"]
                            suffix_str = (
                                str(row["後綴 (Suffix)"]).strip()
                                if pd.notna(row["後綴 (Suffix)"])
                                else ""
                            )
                            has_special_suffix = bool(suffix_str)

                            label_text = build_sample_label(
                                comp,
                                row["劑量 (Dose)"],
                                row["單位 (Unit)"],
                                suffix_str,
                            )
                            try:
                                d_val = float(row["劑量 (Dose)"])
                            except ValueError:
                                d_val = float("inf")

                            sort_weight = (
                                (2, d_val, label_text)
                                if has_special_suffix
                                else (1, d_val, label_text)
                            )

                            if is_superoxide:
                                s_474 = s_diff * 47.4
                                c_474 = (
                                    c_diff * 47.4 if c_diff != 0 else 1.0
                                )
                                rel = (s_474 / c_474) * 100
                                inh = 100 - rel
                                date_drug_rows.append(
                                    (
                                        sort_weight,
                                        [
                                            d,
                                            row["來源檔名"],
                                            label_text,
                                            round(
                                                calc_dict[uid]["min"], 4
                                            ),
                                            round(
                                                calc_dict[uid]["max"], 4
                                            ),
                                            round(s_diff, 4),
                                            round(s_474, 4),
                                            f"{rel:.2f}%"
                                            if not has_special_suffix
                                            else "- (Reference)",
                                            f"{inh:.2f}%"
                                            if not has_special_suffix
                                            else "- (Reference)",
                                        ],
                                    )
                                )
                            else:
                                s_net = s_diff - b_diff
                                c_net = (
                                    c_diff - b_diff
                                    if (c_diff - b_diff) != 0
                                    else 1.0
                                )
                                rel = (s_net / c_net) * 100
                                inh = 100 - rel
                                date_drug_rows.append(
                                    (
                                        sort_weight,
                                        [
                                            d,
                                            row["來源檔名"],
                                            label_text,
                                            round(
                                                calc_dict[uid]["min"], 4
                                            ),
                                            round(
                                                calc_dict[uid]["max"], 4
                                            ),
                                            round(s_diff, 4),
                                            round(s_net, 4),
                                            f"{rel:.2f}%"
                                            if not has_special_suffix
                                            else "- (Reference)",
                                            f"{inh:.2f}%"
                                            if not has_special_suffix
                                            else "- (Reference)",
                                        ],
                                    )
                                )

                        date_drug_rows.sort(key=lambda x: x[0])
                        processed_rows.extend([r[1] for r in date_drug_rows])

                    if clean_sheet_title in wb.sheetnames:
                        ws = wb[clean_sheet_title]
                        for r_vals in processed_rows:
                            ws.append(r_vals)
                    else:
                        ws = wb.create_sheet(title=clean_sheet_title)
                        headers = [
                            "Date",
                            "File Name",
                            "Sample",
                            "Min",
                            "Max",
                            "Max - Min",
                            "*47.4"
                            if is_superoxide
                            else "Max - Min - Basal",
                            "Release (%)",
                            "Inhibition (%)",
                        ]
                        ws.append(
                            [f"Assay: {assay_type}", f"Compound: {comp}", "", ""]
                        )
                        ws.append(
                            [
                                f"Cell: {meta_cell}",
                                f"Stimulant: {meta_stim}",
                                f"Amplifier: {meta_amp}",
                                f"Substrate: {meta_sub}",
                            ]
                        )
                        ws.append([""] * len(headers))
                        ws.append(headers)

                        for r_vals in processed_rows:
                            ws.append(r_vals)

                        for cell in ws[4]:
                            cell.fill = gray_fill
                            cell.font = tnr_bold

                    for row in ws.iter_rows():
                        for cell in row:
                            if (
                                not cell.font
                                or cell.font.name != "Times New Roman"
                            ):
                                cell.font = tnr_font
                            cell.alignment = Alignment(
                                horizontal="center", vertical="center"
                            )

                    for col in ws.columns:
                        max_len = max(
                            len(str(cell.value or "")) for cell in col
                        )
                        col_letter = col[0].column_letter
                        ws.column_dimensions[col_letter].width = max(
                            max_len + 4, 15
                        )

                out_stream = io.BytesIO()
                wb.save(out_stream)

                st.download_button(
                    "下載結構化 Excel 試算表 (.xlsx)",
                    data=out_stream.getvalue(),
                    file_name=out_filename
                    if out_filename.endswith(".xlsx")
                    else f"{out_filename}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                )
                st.success("計算完成並已建立 Excel 工作簿。")
        st.markdown("</div>", unsafe_allow_html=True)

    # -------------------------------------------------------------
    # 分頁 3: 動力學繪圖
    # -------------------------------------------------------------
    with tab_plot:
        plot_mode = st.radio(
            "分析繪圖模式",
            [
                "模式 A：單一藥物單一濃度批量匯出 (原版標準圖)",
                "模式 B：自由組合疊圖分析 (支援差值校正)",
            ],
            horizontal=True,
        )

        ylabel_text = (
            "Absorbance at 550 nm (O.D.)"
            if is_superoxide
            else "Absorbance at 405 nm (O.D.)"
        )
        assay_label = (
            "Superoxide Generation" if is_superoxide else "Elastase Release"
        )

        # ---------------- 模式 A ----------------
        if plot_mode.startswith("模式 A"):
            col_pick_a, col_view_a = st.columns([1.2, 2.8])

            non_ctrl_samples = selected_df[
                ~selected_df["化合物 (Compound)"].str.lower().isin(
                    ["control", "basal"]
                )
            ]
            all_uids_a = non_ctrl_samples["UID"].tolist()

            if "selected_uids_a" not in st.session_state:
                st.session_state.selected_uids_a = all_uids_a.copy()

            with col_pick_a:
                st.markdown(
                    """
                    <div class="workspace-card" style="padding: 12px;">
                        <div class="card-header" style="margin-bottom: 8px;">選擇樣品清單</div>
                    """,
                    unsafe_allow_html=True,
                )

                b_c1, b_c2 = st.columns(2)
                with b_c1:
                    if st.button(
                        "全部選取", use_container_width=True, key="btn_all_a"
                    ):
                        st.session_state.selected_uids_a = all_uids_a.copy()
                        st.rerun()
                with b_c2:
                    if st.button(
                        "全部取消 (清除)",
                        use_container_width=True,
                        key="btn_clr_a",
                    ):
                        st.session_state.selected_uids_a = []
                        st.rerun()

                checked_a = []
                for d in sorted(
                    non_ctrl_samples["日期 (Date)"].unique(),
                    key=lambda x: pd.to_datetime(
                        x, format="%d-%m-%Y", errors="coerce"
                    ),
                ):
                    with st.expander(f"日期：{d}", expanded=True):
                        date_sub = non_ctrl_samples[
                            non_ctrl_samples["日期 (Date)"] == d
                        ]
                        for _, row in date_sub.iterrows():
                            uid = row["UID"]
                            label_fmt = build_sample_label(
                                row["化合物 (Compound)"],
                                row["劑量 (Dose)"],
                                row["單位 (Unit)"],
                                row["後綴 (Suffix)"],
                            )
                            label = f"{label_fmt} ({row['測量時間']})"
                            is_chk = uid in st.session_state.selected_uids_a
                            if st.checkbox(
                                label, value=is_chk, key=f"chk_a_{uid}"
                            ):
                                checked_a.append(uid)

                st.session_state.selected_uids_a = checked_a
                st.caption(f"已選取 {len(checked_a)} 組樣品")
                st.markdown("</div>", unsafe_allow_html=True)

            with col_view_a:
                st.markdown(
                    """
                    <div class="workspace-card">
                        <div class="card-header">即時預覽勾選樣品</div>
                    """,
                    unsafe_allow_html=True,
                )

                if checked_a:
                    for preview_uid in checked_a:
                        p_row = selected_df[
                            selected_df["UID"] == preview_uid
                        ].iloc[0]
                        p_d = p_row["日期 (Date)"]
                        p_comp = p_row["化合物 (Compound)"]
                        p_label = build_sample_label(
                            p_comp,
                            p_row["劑量 (Dose)"],
                            p_row["單位 (Unit)"],
                            p_row["後綴 (Suffix)"],
                        )

                        fig, ax = plt.subplots(figsize=(8.5, 4.8), dpi=150)
                        sources = []

                        if not is_superoxide and p_d in anchor_settings:
                            b_uid = anchor_settings[p_d]["Basal"]
                            b_row = selected_df[
                                selected_df["UID"] == b_uid
                            ].iloc[0]
                            b_df = st.session_state.parsed_data[b_uid]["df"]
                            ax.plot(
                                b_df["Time"],
                                b_df["Abs"],
                                label="Basal (basal)",
                                color="#2ca02c",
                                linewidth=1.8,
                            )
                            sources.append(
                                f" • Basal: {b_row['來源檔名']} ({b_row['測量時間']})"
                            )

                        if p_d in anchor_settings:
                            c_uid = anchor_settings[p_d]["Control"]
                            c_row = selected_df[
                                selected_df["UID"] == c_uid
                            ].iloc[0]
                            c_df = st.session_state.parsed_data[c_uid]["df"]
                            ax.plot(
                                c_df["Time"],
                                c_df["Abs"],
                                label="Control (control)",
                                color="#d62728",
                                linewidth=2.0,
                            )
                            sources.append(
                                f" • Control: {c_row['來源檔名']} ({c_row['測量時間']})"
                            )

                        s_df = st.session_state.parsed_data[preview_uid]["df"]
                        ax.plot(
                            s_df["Time"],
                            s_df["Abs"],
                            label=f"{p_label}",
                            color="#1f77b4",
                            linewidth=2.0,
                        )
                        sources.append(
                            f" • Sample: {p_row['來源檔名']} ({p_row['測量時間']})"
                        )

                        ax.axvline(
                            x=120, color="#94a3b8", linestyle=":", alpha=0.7
                        )
                        ax.axvline(
                            x=240, color="#94a3b8", linestyle=":", alpha=0.7
                        )
                        ax.axvline(
                            x=420, color="#94a3b8", linestyle=":", alpha=0.7
                        )

                        ax.text(
                            120,
                            1.01,
                            "Drug (2')",
                            transform=ax.get_xaxis_transform(),
                            fontsize=8.5,
                            color="#475569",
                            ha="center",
                            va="bottom",
                            fontweight="bold",
                        )
                        ax.text(
                            240,
                            1.01,
                            "CB (4')",
                            transform=ax.get_xaxis_transform(),
                            fontsize=8.5,
                            color="#475569",
                            ha="center",
                            va="bottom",
                            fontweight="bold",
                        )
                        ax.text(
                            420,
                            1.01,
                            "fMLF (7')",
                            transform=ax.get_xaxis_transform(),
                            fontsize=8.5,
                            color="#475569",
                            ha="center",
                            va="bottom",
                            fontweight="bold",
                        )

                        ax.set_xlabel("Time (s)", fontweight="bold")
                        ax.set_ylabel(ylabel_text, fontweight="bold")

                        # 固定 Y 軸範圍與刻度間距
                        ax.set_ylim(y_min, y_max)
                        ax.yaxis.set_major_locator(ticker.MultipleLocator(0.1))
                        ax.yaxis.set_major_formatter(
                            ticker.FormatStrFormatter("%.1f")
                        )

                        ax.set_title(
                            f"{assay_label}: {p_label} ({p_d})",
                            fontsize=11,
                            fontweight="bold",
                            pad=18,
                        )
                        ax.legend(
                            frameon=True,
                            facecolor="white",
                            edgecolor="#cbd5e1",
                            fontsize=8.5,
                            loc="upper left",
                        )
                        ax.grid(True, linestyle="--", alpha=0.3)

                        source_box_str = "Source Files:\n" + "\n".join(sources)
                        line_cnt = len(sources) + 1
                        bottom_margin = max(0.22, 0.10 + 0.035 * line_cnt)
                        plt.subplots_adjust(bottom=bottom_margin)

                        fig.text(
                            0.12,
                            0.02,
                            source_box_str,
                            fontsize=7.5,
                            color="#444444",
                            ha="left",
                            va="bottom",
                            bbox=dict(
                                boxstyle="round,pad=0.3",
                                facecolor="#f9f9f9",
                                edgecolor="#dddddd",
                                alpha=0.8,
                            ),
                        )
                        st.pyplot(fig)
                        plt.close(fig)

                    if st.button(
                        f"批量生成並打包全部 ({len(checked_a)} 張高解析圖表)",
                        use_container_width=True,
                    ):
                        zip_buffer = io.BytesIO()
                        with zipfile.ZipFile(
                            zip_buffer, "w", zipfile.ZIP_DEFLATED
                        ) as zip_file:
                            for uid in checked_a:
                                row = selected_df[
                                    selected_df["UID"] == uid
                                ].iloc[0]
                                d = row["日期 (Date)"]
                                comp = row["化合物 (Compound)"]
                                label_text = build_sample_label(
                                    comp,
                                    row["劑量 (Dose)"],
                                    row["單位 (Unit)"],
                                    row["後綴 (Suffix)"],
                                )

                                fig, ax = plt.subplots(
                                    figsize=(9.5, 6.2), dpi=300
                                )
                                sources = []

                                if not is_superoxide and d in anchor_settings:
                                    b_uid = anchor_settings[d]["Basal"]
                                    b_row = selected_df[
                                        selected_df["UID"] == b_uid
                                    ].iloc[0]
                                    b_df = st.session_state.parsed_data[b_uid][
                                        "df"
                                    ]
                                    ax.plot(
                                        b_df["Time"],
                                        b_df["Abs"],
                                        label="Basal (basal)",
                                        color="#2ca02c",
                                        linewidth=2.0,
                                    )
                                    sources.append(
                                        f" • Basal: {b_row['來源檔名']} ({b_row['測量時間']})"
                                    )

                                if d in anchor_settings:
                                    c_uid = anchor_settings[d]["Control"]
                                    c_row = selected_df[
                                        selected_df["UID"] == c_uid
                                    ].iloc[0]
                                    c_df = st.session_state.parsed_data[c_uid][
                                        "df"
                                    ]
                                    ax.plot(
                                        c_df["Time"],
                                        c_df["Abs"],
                                        label="Control (control)",
                                        color="#d62728",
                                        linewidth=2.2,
                                    )
                                    sources.append(
                                        f" • Control: {c_row['來源檔名']} ({c_row['測量時間']})"
                                    )

                                s_df = st.session_state.parsed_data[uid]["df"]
                                ax.plot(
                                    s_df["Time"],
                                    s_df["Abs"],
                                    label=f"{label_text}",
                                    color="#1f77b4",
                                    linewidth=2.2,
                                )
                                sources.append(
                                    f" • Sample: {row['來源檔名']} ({row['測量時間']})"
                                )

                                ax.axvline(
                                    x=120,
                                    color="#94a3b8",
                                    linestyle=":",
                                    alpha=0.7,
                                )
                                ax.axvline(
                                    x=240,
                                    color="#94a3b8",
                                    linestyle=":",
                                    alpha=0.7,
                                )
                                ax.axvline(
                                    x=420,
                                    color="#94a3b8",
                                    linestyle=":",
                                    alpha=0.7,
                                )

                                ax.text(
                                    120,
                                    1.01,
                                    "Drug (2')",
                                    transform=ax.get_xaxis_transform(),
                                    fontsize=9,
                                    color="#475569",
                                    ha="center",
                                    va="bottom",
                                    fontweight="bold",
                                )
                                ax.text(
                                    240,
                                    1.01,
                                    "CB (4')",
                                    transform=ax.get_xaxis_transform(),
                                    fontsize=9,
                                    color="#475569",
                                    ha="center",
                                    va="bottom",
                                    fontweight="bold",
                                )
                                ax.text(
                                    420,
                                    1.01,
                                    "fMLF (7')",
                                    transform=ax.get_xaxis_transform(),
                                    fontsize=9,
                                    color="#475569",
                                    ha="center",
                                    va="bottom",
                                    fontweight="bold",
                                )

                                ax.set_xlabel(
                                    "Time (s)", fontsize=11, fontweight="bold"
                                )
                                ax.set_ylabel(
                                    ylabel_text, fontsize=11, fontweight="bold"
                                )

                                # 固定 Y 軸範圍與刻度間距
                                ax.set_ylim(y_min, y_max)
                                ax.yaxis.set_major_locator(
                                    ticker.MultipleLocator(0.1)
                                )
                                ax.yaxis.set_major_formatter(
                                    ticker.FormatStrFormatter("%.1f")
                                )

                                ax.set_title(
                                    f"{assay_label}: {label_text} ({d})",
                                    fontsize=13,
                                    fontweight="bold",
                                    pad=18,
                                )
                                ax.grid(True, linestyle="--", alpha=0.5)
                                ax.legend(
                                    frameon=True,
                                    facecolor="white",
                                    edgecolor="#cbd5e1",
                                    framealpha=0.9,
                                    fontsize=9.5,
                                    loc="upper left",
                                )

                                source_box_str = "Source Files:\n" + "\n".join(
                                    sources
                                )
                                line_cnt = len(sources) + 1
                                bottom_margin = max(
                                    0.20, 0.08 + 0.03 * line_cnt
                                )
                                plt.subplots_adjust(bottom=bottom_margin)

                                fig.text(
                                    0.12,
                                    0.02,
                                    source_box_str,
                                    fontsize=8,
                                    color="#444444",
                                    ha="left",
                                    va="bottom",
                                    bbox=dict(
                                        boxstyle="round,pad=0.3",
                                        facecolor="#f9f9f9",
                                        edgecolor="#dddddd",
                                        alpha=0.8,
                                    ),
                                )

                                img_bytes = io.BytesIO()
                                plt.savefig(
                                    img_bytes,
                                    format="png",
                                    dpi=300,
                                    bbox_inches="tight",
                                )
                                plt.close(fig)

                                time_clean = str(row["測量時間"]).replace(
                                    ":", ""
                                )
                                base_fname = os.path.splitext(
                                    row["來源檔名"]
                                )[0]
                                dose_str = (
                                    f"{row['劑量 (Dose)']} {row['單位 (Unit)']}"
                                    if str(row["劑量 (Dose)"]).strip() != "0"
                                    else ""
                                )

                                parts = [
                                    comp,
                                    dose_str,
                                    time_clean,
                                    base_fname,
                                ]
                                img_filename = (
                                    "_".join([p for p in parts if p]).strip()
                                    + ".png"
                                )
                                img_filename = re.sub(
                                    r'[\\/*?:"<>|]', "_", img_filename
                                )
                                zip_file.writestr(
                                    img_filename, img_bytes.getvalue()
                                )

                        st.download_button(
                            "下載所有圖檔打包 ZIP",
                            data=zip_buffer.getvalue(),
                            file_name=f"{assay_label.split()[0]}_Plots_Bundle.zip",
                            mime="application/zip",
                            use_container_width=True,
                        )
                else:
                    st.info("請於左側勾選欲繪圖之樣品。")
                st.markdown("</div>", unsafe_allow_html=True)

        # ---------------- 模式 B ----------------
        else:
            col_pick_b, col_view_b = st.columns([1.3, 2.7])

            all_uids_list = selected_df["UID"].tolist()
            sub_opts = ["None"] + all_uids_list
            names_map = {"None": "無 (原始曲線)"}
            for _, r in selected_df.iterrows():
                label_fmt = build_sample_label(
                    r["化合物 (Compound)"],
                    r["劑量 (Dose)"],
                    r["單位 (Unit)"],
                    r["後綴 (Suffix)"],
                )
                names_map[r["UID"]] = (
                    f"{label_fmt} ({r['日期 (Date)']} {r['測量時間']})"
                )

            selected_plot_pairs = []

            with col_pick_b:
                st.markdown(
                    """
                    <div class="workspace-card" style="padding: 12px;">
                        <div class="card-header" style="margin-bottom: 8px;">自由疊圖清單 (依藥物/日期/濃度)</div>
                    """,
                    unsafe_allow_html=True,
                )

                all_comps_b = sorted(
                    selected_df["化合物 (Compound)"].unique(),
                    key=lambda x: (
                        0
                        if any(
                            k in x.lower()
                            for k in [
                                "basal",
                                "control",
                                "water+water",
                                "water+reagent",
                            ]
                        )
                        else 1,
                        x.lower(),
                    ),
                )

                for comp in all_comps_b:
                    comp_sub = selected_df[
                        selected_df["化合物 (Compound)"] == comp
                    ]
                    comp_uids = comp_sub["UID"].tolist()

                    with st.expander(
                        f"樣品：{comp} ({len(comp_uids)})", expanded=False
                    ):
                        b_group_mode = st.radio(
                            f"分組與快速全選方式 ({comp})",
                            ["依日期分組", "依濃度分組"],
                            horizontal=True,
                            key=f"b_grp_mode_{comp}",
                            label_visibility="collapsed",
                        )

                        if b_group_mode == "依日期分組":
                            comp_dates = sorted(
                                comp_sub["日期 (Date)"].unique(),
                                key=lambda x: pd.to_datetime(
                                    x, format="%d-%m-%Y", errors="coerce"
                                ),
                            )
                            for d_b in comp_dates:
                                date_comp_sub = comp_sub[
                                    comp_sub["日期 (Date)"] == d_b
                                ]

                                t_col, btn_c1, btn_c2 = st.columns(
                                    [2.2, 1.0, 1.0]
                                )
                                with t_col:
                                    st.markdown(
                                        f"<div style='font-size:12px; font-weight:700; color:#0284c7; padding-top:4px;'>📅 {d_b}</div>",
                                        unsafe_allow_html=True,
                                    )
                                with btn_c1:
                                    if st.button(
                                        "全選",
                                        key=f"b_all_d_{comp}_{d_b}",
                                        use_container_width=True,
                                    ):
                                        for u in date_comp_sub["UID"].tolist():
                                            st.session_state[
                                                f"chk_b_comp_{u}"
                                            ] = True
                                        st.rerun()
                                with btn_c2:
                                    if st.button(
                                        "清除",
                                        key=f"b_clr_d_{comp}_{d_b}",
                                        use_container_width=True,
                                    ):
                                        for u in date_comp_sub["UID"].tolist():
                                            st.session_state[
                                                f"chk_b_comp_{u}"
                                            ] = False
                                        st.rerun()

                                for _, r in date_comp_sub.iterrows():
                                    u = r["UID"]
                                    chk_k = f"chk_b_comp_{u}"
                                    r_c1, r_c2 = st.columns([2.5, 1.8])
                                    with r_c1:
                                        lbl = f"{r['樣品名稱 (Sample)']} ({r['測量時間']})"
                                        is_checked = st.checkbox(
                                            lbl,
                                            value=st.session_state.get(
                                                chk_k, False
                                            ),
                                            key=chk_k,
                                        )
                                    with r_c2:
                                        sub_target = st.selectbox(
                                            "扣除",
                                            sub_opts,
                                            index=0,
                                            format_func=lambda x: names_map.get(
                                                x, x
                                            ),
                                            key=f"sub_sel_{u}",
                                            label_visibility="collapsed",
                                        )
                                    if is_checked:
                                        selected_plot_pairs.append(
                                            (u, sub_target)
                                        )

                        else:
                            comp_doses = sorted(
                                comp_sub["劑量 (Dose)"].unique(),
                                key=lambda x: float(x)
                                if str(x).replace(".", "", 1).isdigit()
                                else 9999,
                            )
                            for dose_val in comp_doses:
                                dose_unit = comp_sub[
                                    comp_sub["劑量 (Dose)"] == dose_val
                                ]["單位 (Unit)"].iloc[0]
                                dose_display = (
                                    f"{dose_val} {dose_unit}"
                                    if str(dose_val) != "0"
                                    else "0"
                                )
                                dose_comp_sub = comp_sub[
                                    comp_sub["劑量 (Dose)"] == dose_val
                                ]

                                t_col, btn_c1, btn_c2 = st.columns(
                                    [2.2, 1.0, 1.0]
                                )
                                with t_col:
                                    st.markdown(
                                        f"<div style='font-size:12px; font-weight:700; color:#059669; padding-top:4px;'>🧪 濃度: {dose_display}</div>",
                                        unsafe_allow_html=True,
                                    )
                                with btn_c1:
                                    if st.button(
                                        "全選",
                                        key=f"b_all_dose_{comp}_{dose_val}",
                                        use_container_width=True,
                                    ):
                                        for u in dose_comp_sub["UID"].tolist():
                                            st.session_state[
                                                f"chk_b_comp_{u}"
                                            ] = True
                                        st.rerun()
                                with btn_c2:
                                    if st.button(
                                        "清除",
                                        key=f"b_clr_dose_{comp}_{dose_val}",
                                        use_container_width=True,
                                    ):
                                        for u in dose_comp_sub["UID"].tolist():
                                            st.session_state[
                                                f"chk_b_comp_{u}"
                                            ] = False
                                        st.rerun()

                                for _, r in dose_comp_sub.iterrows():
                                    u = r["UID"]
                                    chk_k = f"chk_b_comp_{u}"
                                    r_c1, r_c2 = st.columns([2.5, 1.8])
                                    with r_c1:
                                        lbl = f"{r['日期 (Date)']} ({r['測量時間']})"
                                        is_checked = st.checkbox(
                                            lbl,
                                            value=st.session_state.get(
                                                chk_k, False
                                            ),
                                            key=chk_k,
                                        )
                                    with r_c2:
                                        sub_target = st.selectbox(
                                            "扣除",
                                            sub_opts,
                                            index=0,
                                            format_func=lambda x: names_map.get(
                                                x, x
                                            ),
                                            key=f"sub_sel_{u}",
                                            label_visibility="collapsed",
                                        )
                                    if is_checked:
                                        selected_plot_pairs.append(
                                            (u, sub_target)
                                        )

                st.caption(f"已選取 {len(selected_plot_pairs)} 條曲線")
                st.markdown("</div>", unsafe_allow_html=True)

            with col_view_b:
                st.markdown(
                    """
                    <div class="workspace-card">
                        <div class="card-header">疊圖即時分析</div>
                    """,
                    unsafe_allow_html=True,
                )

                if selected_plot_pairs:
                    fig, ax = plt.subplots(figsize=(9, 5.5), dpi=150)
                    high_contrast_colors = [
                        "#1f77b4",
                        "#d62728",
                        "#2ca02c",
                        "#9467bd",
                        "#ff7f0e",
                        "#8c564b",
                        "#e377c2",
                        "#17becf",
                        "#bcbd22",
                        "#393b79",
                        "#637939",
                        "#8c6d31",
                        "#843c39",
                        "#7b4173",
                        "#3182bd",
                        "#31a354",
                        "#e6550d",
                        "#756bb1",
                        "#005a32",
                        "#99000d",
                    ]
                    colors = [
                        high_contrast_colors[i % len(high_contrast_colors)]
                        for i in range(len(selected_plot_pairs))
                    ]

                    source_texts = []
                    label_info_list = []

                    for idx, (uid_main, uid_sub) in enumerate(
                        selected_plot_pairs
                    ):
                        row_main = selected_df[
                            selected_df["UID"] == uid_main
                        ].iloc[0]
                        df_main = st.session_state.parsed_data[uid_main][
                            "df"
                        ].copy()

                        label_main_fmt = build_sample_label(
                            row_main["化合物 (Compound)"],
                            row_main["劑量 (Dose)"],
                            row_main["單位 (Unit)"],
                            row_main["後綴 (Suffix)"],
                        )
                        label_main = (
                            f"{label_main_fmt} ({row_main['日期 (Date)']})"
                        )

                        if (
                            uid_sub != "None"
                            and uid_sub in st.session_state.parsed_data
                        ):
                            df_sub = st.session_state.parsed_data[uid_sub][
                                "df"
                            ].copy()
                            row_sub = selected_df[
                                selected_df["UID"] == uid_sub
                            ].iloc[0]
                            merged_df = pd.merge_asof(
                                df_main.sort_values("Time"),
                                df_sub.sort_values("Time"),
                                on="Time",
                                suffixes=("_A", "_B"),
                                direction="nearest",
                            )
                            plot_times = merged_df["Time"]
                            plot_values = (
                                merged_df["Abs_A"] - merged_df["Abs_B"]
                            )

                            label_sub_simple = build_sample_label(
                                row_sub["化合物 (Compound)"],
                                row_sub["劑量 (Dose)"],
                                row_sub["單位 (Unit)"],
                                row_sub["後綴 (Suffix)"],
                            )
                            final_label = (
                                f"{label_main} - {label_sub_simple}"
                            )
                            source_texts.append(
                                f"• {final_label}: {row_main['來源檔名']} - {row_sub['來源檔名']}"
                            )
                        else:
                            plot_times = df_main["Time"]
                            plot_values = df_main["Abs"]
                            final_label = label_main
                            source_texts.append(
                                f"• {final_label}: {row_main['來源檔名']}"
                                f" ({row_main['測量時間']})"
                            )

                        ax.plot(
                            plot_times,
                            plot_values,
                            "-",
                            linewidth=2.0,
                            color=colors[idx],
                            label=final_label,
                        )
                        label_info_list.append(
                            {
                                "target_x": plot_times.iloc[-1],
                                "target_y": plot_values.iloc[-1],
                                "current_y": plot_values.iloc[-1],
                                "text": final_label,
                                "color": colors[idx],
                            }
                        )

                    if label_info_list:
                        label_info_list.sort(key=lambda item: item["target_y"])
                        y_bottom_val, y_top_val = ax.get_ylim()
                        min_safe_gap = (y_top_val - y_bottom_val) * 0.038
                        for i in range(1, len(label_info_list)):
                            prev_y = label_info_list[i - 1]["current_y"]
                            curr_y = label_info_list[i]["current_y"]
                            if curr_y - prev_y < min_safe_gap:
                                label_info_list[i]["current_y"] = (
                                    prev_y + min_safe_gap
                                )

                        for item in label_info_list:
                            ax.text(
                                item["target_x"] + 5,
                                item["current_y"],
                                item["text"],
                                color=item["color"],
                                va="center",
                                fontweight="bold",
                                fontsize=8,
                            )

                    ax.axvline(
                        x=120, color="#94a3b8", linestyle=":", alpha=0.7
                    )
                    ax.axvline(
                        x=240, color="#94a3b8", linestyle=":", alpha=0.7
                    )
                    ax.axvline(
                        x=420, color="#94a3b8", linestyle=":", alpha=0.7
                    )

                    ax.text(
                        120,
                        1.01,
                        "Drug (2')",
                        transform=ax.get_xaxis_transform(),
                        fontsize=8.5,
                        color="#475569",
                        ha="center",
                        va="bottom",
                        fontweight="bold",
                    )
                    ax.text(
                        240,
                        1.01,
                        "CB (4')",
                        transform=ax.get_xaxis_transform(),
                        fontsize=8.5,
                        color="#475569",
                        ha="center",
                        va="bottom",
                        fontweight="bold",
                    )
                    ax.text(
                        420,
                        1.01,
                        "fMLF (7')",
                        transform=ax.get_xaxis_transform(),
                        fontsize=8.5,
                        color="#475569",
                        ha="center",
                        va="bottom",
                        fontweight="bold",
                    )

                    ax.set_xlabel("Time (s)", fontweight="bold", fontsize=10)
                    ax.set_ylabel(ylabel_text, fontweight="bold", fontsize=10)

                    # 固定 Y 軸範圍與刻度間距
                    ax.set_ylim(y_min, y_max)
                    ax.yaxis.set_major_locator(ticker.MultipleLocator(0.1))
                    ax.yaxis.set_major_formatter(
                        ticker.FormatStrFormatter("%.1f")
                    )

                    ax.set_title(
                        f"{assay_label} - Kinetics Overlay",
                        fontweight="bold",
                        fontsize=12,
                        pad=18,
                    )
                    ax.spines["top"].set_visible(False)
                    ax.spines["right"].set_visible(False)
                    ax.grid(True, linestyle="--", alpha=0.4)

                    source_cnt = len(source_texts) + 1
                    bot_padding = max(
                        0.20, min(0.48, 0.08 + 0.025 * source_cnt)
                    )
                    plt.subplots_adjust(bottom=bot_padding)

                    fig.text(
                        0.12,
                        0.02,
                        "Source Files:\n" + "\n".join(source_texts),
                        fontsize=7.5,
                        color="#444444",
                        ha="left",
                        va="bottom",
                        bbox=dict(
                            facecolor="#f9f9f9",
                            edgecolor="#dddddd",
                            alpha=0.8,
                            boxstyle="round,pad=0.3",
                        ),
                    )

                    st.pyplot(fig)

                    img_buf = io.BytesIO()
                    fig.savefig(
                        img_buf, format="png", bbox_inches="tight", dpi=300
                    )
                    plt.close(fig)

                    now_str = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
                    dl_col1, dl_col2 = st.columns([2.5, 1.5])
                    with dl_col1:
                        custom_b_name = st.text_input(
                            "輸出檔名:",
                            value=f"Overlay_{now_str}.png",
                            key="name_input_mode_b",
                        )
                    with dl_col2:
                        st.write("")
                        st.write("")
                        out_b_filename = (
                            custom_b_name.strip()
                            if custom_b_name.strip().endswith(".png")
                            else f"{custom_b_name.strip()}.png"
                        )
                        st.download_button(
                            "下載疊圖 (PNG)",
                            data=img_buf.getvalue(),
                            file_name=out_b_filename,
                            mime="image/png",
                            use_container_width=True,
                        )
                else:
                    st.info("請於左側選擇欲疊合之曲線。")
                st.markdown("</div>", unsafe_allow_html=True)
